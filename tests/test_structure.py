"""The requested document structure must survive every rendering path."""
import copy
import json
from pathlib import Path

from docx import Document
import fitz

from architect_agent.schemas import ARCHITECTURE_STRUCTURE, UPDATE_SCHEMA, validate
from architect_agent.output import report_sections, report_html, report_pdf
from architect_agent.notion import blocks_for_plan
from architect_agent.workflow import build
from architect_agent.sources import Source

TITLES = [
    '1. Managementsamenvatting', '2. Opdracht, scope en bronnen',
    '3. Klantroutes en dienstverlening', '4. Bestaande en benodigde bedrijfsservices',
    '5. Oplossing in vijf lagen', '6. Uitwerking van nieuwe of gewijzigde bedrijfsservices',
    '7. Proces- en gegevensstromen', '8. Informatie en registers',
    '9. Architectuurbesluiten en beslisregels', '10. Privacy, beveiliging en kwaliteit',
    '11. Realisatie, deployment en beheer', '12. Acceptatie en open besluiten',
]
ROOT = Path(__file__).resolve().parents[1]


def example():
    return json.loads((ROOT/'examples/example_plan.json').read_text())


def sources():
    return [Source(f'B{i:02}', f'bron{i}.pdf', 'Niet publiceren: volledige broninhoud.', 30, 'pdf') for i in range(3)]


def test_twelve_chapters_preserve_legacy_content_and_mark_gaps():
    plan = example()
    assert [chapter['title'] for chapter in ARCHITECTURE_STRUCTURE] == TITLES
    sections = report_sections(plan, sources())
    assert [title for title, lines in sections] == TITLES
    text = '\n'.join(line for title, lines in sections for line in lines)
    assert plan['management_summary'] in text
    assert all(change['change'] in text for change in plan['changes'])
    assert 'Nog te bepalen' in sections[2][1][0]
    assert 'Nog te bepalen' in sections[7][1][0]
    assert all(s.text not in text for s in sources())


def test_new_chapter_content_is_validated_and_preserved_in_exports(tmp_path):
    plan = example()
    plan['chapter_content'] = {chapter['key']:[{'text': 'Uitwerking '+chapter['key'], 'evidence':['B01 p.7']}] for chapter in ARCHITECTURE_STRUCTURE}
    assert not validate(plan, sources())
    assert 'chapter_content' in UPDATE_SCHEMA['required']
    html = report_html(plan, sources(), 'Rotterdam')
    assert all(html.index(TITLES[i]) < html.index(TITLES[i+1]) for i in range(11))
    assert 'Uitwerking customer_routes' in html
    assert 'Bron: B01 p.7' in html
    blocks = blocks_for_plan(plan, sources())
    headings = [''.join(r['text']['content'] for r in b['heading_2']['rich_text']) for b in blocks if b['type']=='heading_2']
    assert headings == TITLES
    pdf = report_pdf(plan, sources(), 'Rotterdam')
    with fitz.open(stream=pdf, filetype='pdf') as doc:
        text = ' '.join(' '.join(p.get_text().split()) for p in doc)
    assert all(title in text for title in TITLES)
    base = tmp_path/'base.docx'
    original = Document(); original.add_paragraph('Bestaande inhoud behouden'); original.save(base)
    output = tmp_path/'output'
    manifest = build(base, sources(), plan, output)
    doc = Document(output/manifest['document'])
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 2'] == TITLES
    assert 'Bestaande inhoud behouden' in [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 2
    assert any(p.text == 'Uitwerking information\nBron: B01 p.7' for p in doc.paragraphs)


def test_incomplete_chapters_and_unknown_references_fail_validation():
    plan = example()
    plan['chapter_content'] = {c['key']:[{'text':'Nog te bepalen', 'evidence':[]}] for c in ARCHITECTURE_STRUCTURE}
    broken = copy.deepcopy(plan); del broken['chapter_content']['customer_routes']
    assert any('twaalf' in error for error in validate(broken, sources()))
    broken = copy.deepcopy(plan); broken['chapter_content']['customer_routes'] = []
    assert any('customer_routes' in error for error in validate(broken, sources()))
    plan['chapter_content']['customer_routes'][0]['evidence'] = ['B99 p.1']
    assert any('Bron ontbreekt' in error for error in validate(plan, sources()))
