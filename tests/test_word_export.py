"""Validate the local browser DOCX exporter using a real Word reader."""
import json
from pathlib import Path
import shutil
import subprocess
from zipfile import ZipFile

from docx import Document
import pytest

from architect_agent.output import STYLES, report_sections
from architect_agent.schemas import ARCHITECTURE_STRUCTURE
from architect_agent.sources import Source

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.skipif(shutil.which('node') is None, reason='Browserexportcontrole vereist Node.js.')
@pytest.mark.parametrize('brand', ['Rotterdam', 'Common Ground'])
def test_browser_docx_preserves_every_section_and_unicode(tmp_path, brand):
    plan = json.loads((ROOT/'examples/example_plan.json').read_text())
    sources = [Source(f'B{i:02}',f'bron{i}.pdf','Volledige bron niet publiceren',30,'pdf') for i in range(3)]
    sections = report_sections(plan, sources)
    sections[2][1].append('Letterlijke tekst: <voorstel> & “klant” — één.')
    target = tmp_path/'architectuur.docx'
    script = '''
const fs=require('fs'),vm=require('vm');
vm.runInThisContext(fs.readFileSync('assets/vendor/docx/docx.iife.js','utf8'));
vm.runInThisContext(fs.readFileSync('web/word_export.js','utf8'));
(async()=>{const data=JSON.parse(fs.readFileSync(0,'utf8'));
 const blob=await WordExport.create(data.title,data.sections,data.brand,data.style);
 fs.writeFileSync(process.argv[1],Buffer.from(await blob.arrayBuffer()));
})().catch(e=>{console.error(e);process.exit(1)});
'''
    subprocess.run(['node','-e',script,str(target)],cwd=ROOT,input=json.dumps({
        'title':plan['title'],'sections':sections,'brand':brand,'style':STYLES[brand]
    }),text=True,check=True,capture_output=True)
    with ZipFile(target) as archive:
        assert archive.testzip() is None
        assert 'word/document.xml' in archive.namelist()
    doc = Document(target)
    assert doc.paragraphs[0].text == plan['title']
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 1'] == [c['title'] for c in ARCHITECTURE_STRUCTURE]
    text='\n'.join(p.text for p in doc.paragraphs)
    assert all(line in text for title,lines in sections for line in lines)
    assert 'Letterlijke tekst: <voorstel> & “klant” — één.' in text
    assert all(source.text not in text for source in sources)
    assert brand in text and brand in doc.sections[0].footer.paragraphs[0].text
    assert str(next(p.style for p in doc.paragraphs if p.style.name == 'Heading 1').font.color.rgb) == STYLES[brand]['primary'][1:]
