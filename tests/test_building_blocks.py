import json
from pathlib import Path
import subprocess
import shutil
from unittest.mock import patch

from docx import Document
import fitz
from streamlit.testing.v1 import AppTest
import pytest

from architect_agent.building_blocks import (
    advice_document, catalogue_from_snapshot, load_catalogue, recommend, selection_manifest,
)
from architect_agent.selection import SelectedDocument

ROOT = Path(__file__).parents[1]


def test_every_mapping_has_exact_source_quote_and_no_invented_product():
    catalogue = load_catalogue()
    assert len(catalogue['abbs']) == 34
    assert len(catalogue['sbbs']) == 30
    source = catalogue['source']
    assert source['sha256'] and source['url'].endswith('Enterprisearchitectuur-GPT.md')
    for abb in catalogue['abbs']:
        assert abb['name'] in abb['quote']
        assert abb['line'] > 0
        for sbb_id in abb['sbb_ids']:
            sbb = next(b for b in catalogue['sbbs'] if b['id'] == sbb_id)
            assert sbb['name'] in abb['quote']
            assert abb['id'] in sbb['abb_ids']


def test_catalogue_parser_keeps_scope_and_layer():
    source = {'text': '## Solution Building Blocks (SBB’s) en mapping op ABB’s\n\n#### Laag 5 – Interactie\n| **ABB (service die…)** | **Primaire SBB** |\n|---|---|\n| Digitale formulieren aanbiedt en verwerkt | Open Formulieren |\n## Ander hoofdstuk\n| Dit hoort niet bij de mapping | Fantasieproduct |\n',
              'url': 'https://example.test/ea', 'sha256': 'abc', 'retrieved_at': '2026-10-16'}
    catalogue = catalogue_from_snapshot(source)
    assert len(catalogue['abbs']) == len(catalogue['sbbs']) == 1
    assert catalogue['sbbs'][0]['name'] == 'Open Formulieren'
    assert catalogue['abbs'][0]['layer'] == 'Laag 5 – Interactie'


def test_recommendations_follow_topic_and_contents_with_evidence():
    catalogue = load_catalogue()
    contexts = [{'source': 'Onderwerp', 'text': 'Zaakgericht werken'},
                {'source': 'project.md', 'text': 'We bieden digitale formulieren met notificaties na de aanvraag.'}]
    advice = recommend(catalogue, contexts)
    manifest = selection_manifest(catalogue, advice, {})
    names = {b['name'] for b in manifest['selected'] if b['type'] == 'SBB'}
    assert {'OpenZaak API', 'OpenZaak Registratie', 'Open Formulieren', 'OpenNotificaties'} <= names
    assert 'Keycloak / Autorisatieservice' not in names
    assert any(e['source'] == 'project.md' and 'notificaties' in e['excerpt']
               for item in advice for e in item['evidence'])
    assert not recommend(catalogue, [{'source': 'Onderwerp', 'text': ''}])


def test_auto_selection_changes_and_manual_override_survives():
    catalogue = load_catalogue()
    forms = recommend(catalogue, [{'source': 'Onderwerp', 'text': 'digitale formulieren'}])
    selected = selection_manifest(catalogue, forms, {})
    form_abb = next(b for b in selected['selected'] if b['type'] == 'ABB')
    excluded = selection_manifest(catalogue, forms, {form_abb['id']: False})
    assert not excluded['selected']
    changed = recommend(catalogue, [{'source': 'Onderwerp', 'text': 'Autorisatie en authenticatie'}])
    assert all(b['name'] != 'Open Formulieren' for b in selection_manifest(catalogue, changed, {})['selected'])
    kept = selection_manifest(catalogue, changed, {form_abb['id']: True})
    assert any(b['name'] == 'Open Formulieren' for b in kept['selected'])
    document = advice_document(kept, catalogue)
    assert 'VOORSTEL' in document and form_abb['name'] in document
    assert catalogue['source']['sha256'] in document


@pytest.mark.skipif(shutil.which('node') is None, reason='Optionele JavaScript/Python-pariteitscontrole vereist Node.js.')
def test_javascript_and_python_agree_on_advice_and_overrides():
    catalogue = load_catalogue()
    contexts = [{'source': 'Onderwerp', 'text': 'Zaken, formulieren en privacy'},
                {'source': 'eisen.pdf p.2', 'text': 'Notificaties, logging en autorisatie zijn nodig.'}]
    advice = recommend(catalogue, contexts)
    overrides = {advice[0]['id']: False}
    script = "const fs=require('fs');const rules=require('./web/building_blocks.js');const input=JSON.parse(fs.readFileSync(0,'utf8'));const advice=rules.recommend(input.catalogue,input.contexts);process.stdout.write(JSON.stringify(rules.selection(input.catalogue,advice,input.overrides)));"
    response = subprocess.run(['node', '-e', script], cwd=ROOT,
                              input=json.dumps({'catalogue': catalogue, 'contexts': contexts, 'overrides': overrides}),
                              text=True, capture_output=True, check=True)
    assert json.loads(response.stdout) == selection_manifest(catalogue, advice, overrides)


def test_python_upload_reader_uses_pdf_docx_and_deselected_docs(tmp_path):
    from architect_agent.building_blocks_ui import _contexts
    pdf_path = tmp_path/'bron.pdf'
    with fitz.open() as pdf:
        page = pdf.new_page();page.insert_text((72, 72), 'Notificaties worden verzonden na het verwerken van een aanvraag.')
        pdf.save(pdf_path)
    docx_path = tmp_path/'bron.docx'
    doc = Document();doc.add_paragraph('Authenticatie en autorisatie verlopen via Keycloak voor deze voorziening.');doc.save(docx_path)
    docs = [SelectedDocument('project', pdf_path.name, pdf_path.read_bytes()),
            SelectedDocument('project', docx_path.name, docx_path.read_bytes()),
            SelectedDocument('enterprise', 'ea.md', b'Dit algemene kader noemt alle bouwblokken, logging en privacy.')]
    # De functie werkt lokaal; er wordt geen model of netwerk aangeroepen.
    with patch('architect_agent.building_blocks_ui.st.session_state', {}):
        contexts, warnings = _contexts('Voorziening', '', docs, None)
    assert not warnings
    assert any('Notificaties' in c['text'] for c in contexts)
    assert any('Keycloak' in c['text'] for c in contexts)
    assert all('algemene kader' not in c['text'] for c in contexts)


def test_streamlit_auto_choices_manual_reset_and_psa_source():
    at = AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
    at.text_input(key='project_name').set_value('Digitale formulieren').run()
    assert not at.exception
    catalogue = load_catalogue()
    form = next(b for b in catalogue['abbs'] if b['name'] == 'Digitale formulieren aanbiedt en verwerkt')
    choice = 'building_block_'+form['id']
    assert at.checkbox(key=choice).value
    at.checkbox(key=choice).uncheck().run()
    assert not at.checkbox(key=choice).value
    at.text_area(key='project_goals').set_value('Ook digitale formulieren voor aanvragen.').run()
    assert not at.checkbox(key=choice).value
    at.button(key='reset_building_blocks').click().run()
    assert at.checkbox(key=choice).value
    # Advies wordt daadwerkelijk aan de bronset toegevoegd voor conceptgeneratie.
    assert any('Bouwblokkenadvies.md' in frame.value['Bron'].tolist() for frame in at.dataframe)
    at.text_input(key='project_name').set_value('Autorisatie').run()
    at.text_area(key='project_goals').set_value('Toegang via authenticatie.').run()
    assert not at.checkbox(key=choice).value
    assert not at.exception
