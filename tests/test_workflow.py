from pathlib import Path
import json
import pytest
from docx import Document
from architect_agent.sources import collect
from architect_agent.schemas import validate
from architect_agent.workflow import build
from architect_agent.word import START,END

ROOT=Path(__file__).resolve().parents[1]


def fixture_files(tmp_path):
    p=tmp_path/'basis.docx';d=Document();d.add_heading('Solutionarchitectuur',0);d.add_paragraph('Oorspronkelijke tekst mag niet verdwijnen.');d.save(p)
    f1=tmp_path/'signaal.md'; f1.write_text('Signalen blijven operationeel in bronsystemen. Een nieuw datawarehouse voert tactische analyse uit.\n'*12)
    f2=tmp_path/'adr.md'; f2.write_text('Bestaande ADR-009: klanttevredenheidscomponent niet prioriteren. Concept-aanpak toegestaan.\n'*12)
    return p,[f1,f2]


def demo_plan():
    return json.loads((ROOT/'examples/example_plan.json').read_text(encoding='utf-8'))


def test_schema_and_sources(tmp_path):
    base,paths=fixture_files(tmp_path)
    parsed=collect(paths,base_doc=base)
    assert [s.key for s in parsed]==['B00','B01','B02']
    # De demo gebruikt PDF-paginakenmerken. Voor synthetische Markdown-bronnen worden ze als p.1 gelezen.
    plan=demo_plan()
    for ch in plan['changes']: ch['evidence']=['B01 tekst']
    for rule in plan['decision_rules']: rule['evidence']=['B02 tekst']
    assert not validate(plan,parsed)
    plan['adr']['status']='Vastgesteld'
    assert any('CONCEPT' in s for s in validate(plan,parsed))


def test_repeat_replaces_only_previous_addition(tmp_path):
    base,paths=fixture_files(tmp_path)
    parsed=collect(paths,base_doc=base)
    plan=demo_plan()
    for ch in plan['changes']: ch['evidence']=['B01 tekst']
    for rule in plan['decision_rules']: rule['evidence']=['B02 tekst']
    out=tmp_path/'run1';build(base,parsed,plan,out)
    doc=Document(out/'Solutionarchitectuur_bijgewerkt_concept.docx')
    full='\n'.join(p.text for p in doc.paragraphs)
    assert 'Oorspronkelijke tekst mag niet verdwijnen' in full
    assert START in full and END in full
    assert len(doc.inline_shapes)==2
    out2=tmp_path/'run2';build(out/'Solutionarchitectuur_bijgewerkt_concept.docx',parsed,plan,out2)
    doc2=Document(out2/'Solutionarchitectuur_bijgewerkt_concept.docx')
    full2='\n'.join(p.text for p in doc2.paragraphs)
    assert full2.count(START)==1
    assert full2.count(END)==1
    assert full2.count('Architectuurupdate — concept ter review')==1
    assert len(doc2.inline_shapes)==2
    assert (out2/'architectuur_vijflagen.svg').exists()
    assert (out2/'gegevensstroom.png').exists()


def test_bad_page_citation_rejected(tmp_path):
    base,paths=fixture_files(tmp_path)
    parsed=collect(paths,base_doc=base)
    plan=demo_plan()
    plan['changes'][0]['evidence']=['B77 p.42']
    assert any('Bron ontbreekt' in s for s in validate(plan,parsed))
