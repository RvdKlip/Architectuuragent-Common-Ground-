"""Niet-destructieve update van een bestaand Word-document; herhaalbaar via markers."""
from __future__ import annotations
from datetime import date
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

START = '[[CG_ARCH_AGENT_UPDATE_START]]'
END = '[[CG_ARCH_AGENT_UPDATE_END]]'


def remove_previous_agent_update(doc: Document) -> bool:
    body=doc.element.body
    children=list(body)
    start=None; end=None
    for idx,child in enumerate(children):
        text=''.join(child.itertext())
        if START in text and start is None: start=idx
        if END in text: end=idx
    if start is None:
        return False
    if end is None or end < start:
        raise ValueError('Een vorige agent-update heeft geen geldige eindmarker. Bestand niet gewijzigd.')
    for child in children[start:end+1]:
        body.remove(child)
    return True


def theme(doc):
    styles=doc.styles
    normal=styles['Normal']
    normal.font.name='Aptos'
    normal.font.size=Pt(9)
    for key,size in [('Heading 1',16),('Heading 2',12),('Heading 3',10)]:
        style=styles[key]
        style.font.name='Aptos Display'
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(13,79,103)


def marker(doc, text):
    p=doc.add_paragraph(style='Normal')
    run=p.add_run(text)
    run.font.hidden=True
    run.font.size=Pt(1)


def para(doc, text, style=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.add_run(text)
    return p


def bullets(doc, items):
    for item in items:
        para(doc,str(item), 'List Bullet')


def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=True
    for i,name in enumerate(headers):
        c=t.rows[0].cells[i]
        c.text=name
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs:
                r.bold=True
                r.font.size=Pt(8)
                r.font.color.rgb=RGBColor(255,255,255)
        tcPr=c._tc.get_or_add_tcPr()
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'0D4F67'); tcPr.append(shd)
    header=t.rows[0]
    trPr=header._tr.get_or_add_trPr()
    hd=OxmlElement('w:tblHeader'); hd.set(qn('w:val'),'true'); trPr.append(hd)
    for values in rows:
        cells=t.add_row().cells
        for i,val in enumerate(values):
            cells[i].text=str(val)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(8)
        rowPr=t.rows[-1]._tr.get_or_add_trPr()
        rule=OxmlElement('w:cantSplit'); rowPr.append(rule)
    doc.add_paragraph()
    return t


def image(doc, location, caption):
    sec=doc.sections[-1]
    width=sec.page_width-sec.left_margin-sec.right_margin
    max_width=min(float(width)/360000, 16.6)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(location), width=Cm(max_width))
    cp=doc.add_paragraph()
    cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=cp.add_run(caption)
    r.italic=True; r.font.size=Pt(8)


def make_document(base_doc: Path, plan: dict, sources: list, figures: dict, destination: Path):
    doc=Document(str(base_doc))
    remove_previous_agent_update(doc)
    # Bestaande styles en tekst blijven ongemoeid. Nieuwe tekst gebruikt bestaande Heading-stijlen.
    marker(doc, START)
    doc.add_page_break()
    doc.add_heading('Architectuurupdate — concept ter review', 1)
    para(doc, f"Gegenereerd op {date.today().strftime('%d-%m-%Y')} · onderwerp: {plan['scope']}")
    para(doc, 'Deze gecontroleerde aanvulling wijzigt de bestaande hoofdstukken niet. De inhoud is een voorstel, geen vastgesteld architectuurbesluit. Bij volgende runs vervangt de agent uitsluitend deze eigen aanvulling.')

    from .output import report_sections
    for index, (title, lines) in enumerate(report_sections(plan, sources), 1):
        doc.add_heading(title, 2)
        if index == 5:
            image(doc, figures['layers'], 'Positionering binnen het vijflagenmodel van Platform Dienstverlening — concept.')
        elif index == 7:
            image(doc, figures['flow'], 'Voorgestelde proces- en gegevensstroom — concept; toetsen aan bron- en privacyafspraken.')
        for line in lines:
            para(doc, str(line))
    marker(doc, END)
    destination=Path(destination)
    destination.parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(destination))
    return destination
