"""Lokale, controleerbare bronextractie. Geen OCR en geen externe verbindingen."""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
import re

import fitz
from docx import Document

SOURCE_LIMIT = 900_000

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class Source:
    key: str
    filename: str
    text: str
    pages: int
    kind: str

    @property
    def reference(self) -> str:
        return f"{self.key} — {self.filename}"


def extract(path: Path, key: str) -> Source:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Niet-ondersteund bestand: {path.name} (gebruik PDF, Word, TXT of Markdown)")
    sections: list[str] = []
    n_pages = 0
    if suffix == '.pdf':
        with fitz.open(str(path)) as pdf:
            n_pages = len(pdf)
            visual_only_pages = []
            for page_index, page in enumerate(pdf, 1):
                contents = page.get_text('text', sort=True).strip()
                if contents:
                    sections.append(f"[{key} p.{page_index}]\n{contents}")
                elif page.get_images(full=False):
                    visual_only_pages.append(page_index)
            if not sections:
                raise ValueError(f"PDF bevat geen uitleesbare tekst: {path.name}. Lever een tekst-PDF aan.")
            if visual_only_pages:
                sections.append(f"[{key} WAARSCHUWING: pagina's {visual_only_pages} bevatten afbeeldingen zonder uitleesbare tekst en zijn niet inhoudelijk geanalyseerd. Raadpleeg ze handmatig.]")
    elif suffix == '.docx':
        doc = Document(str(path))
        parts = []
        in_agent_update = False
        for paragraph in doc.paragraphs:
            value = paragraph.text.strip()
            if '[[CG_ARCH_AGENT_UPDATE_START]]' in value:
                in_agent_update = True
                continue
            if '[[CG_ARCH_AGENT_UPDATE_END]]' in value:
                in_agent_update = False
                continue
            if value and not in_agent_update:
                parts.append(value)
        for i, t in enumerate(doc.tables, 1):
            parts.append(f"[Tabel {i}]\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows))
        sections = [f"[{key} doc]\n" + '\n'.join(parts)]
        n_pages = 1
    else:
        sections = [f"[{key} tekst]\n" + path.read_text(encoding='utf-8-sig')]
        n_pages = 1
    cleaned = '\n\n'.join(sections).strip()
    if len(re.sub(r'\W', '', cleaned)) < 30:
        raise ValueError(f"Er is vrijwel geen tekst gevonden in {path.name}.")
    return Source(key, path.name, cleaned, n_pages, suffix[1:])


def collect(paths: list[Path], *, base_doc: Path | None = None, limit: int = SOURCE_LIMIT) -> list[Source]:
    all_paths = ([Path(base_doc)] if base_doc else []) + [Path(p) for p in paths]
    result = [extract(p, f"B{i:02d}") for i, p in enumerate(all_paths)]
    size = sum(len(s.text) for s in result)
    if size > limit:
        raise ValueError(
            f"De bronset bevat {size:,} tekens en overschrijdt de limiet van {limit:,}. "
            "Verklein de bronset; teksten worden nooit stilzwijgend afgekapt."
        )
    return result


def source_context(sources: list[Source]) -> str:
    return "\n\n".join("=" * 25 + f" BRON {s.reference} " + "=" * 25 + '\n' + s.text for s in sources)
